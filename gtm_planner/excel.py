"""Auditable Excel model: inputs in blue, formulas for hours/capacity, engine reconciliation."""
from datetime import date
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HFONT = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="14213D")
INPUT = Font(color="0000FF")
GOOD = PatternFill("solid", fgColor="D8F3DC")
BAD = PatternFill("solid", fgColor="F8D7DA")
WARN = PatternFill("solid", fgColor="FFF3CD")


def _hdr(ws, row, values, col=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.font, c.fill = HFONT, HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(path, ctx):
    a, clean, results, rec = ctx["a"], ctx["clean"], ctx["results"], ctx["rec"]
    plan, avail, demand, sens = ctx["plan"], ctx["avail"], ctx["demand"], ctx["sens"]
    teams, metrics = ctx["teams"], results[rec]["metrics"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    dq, ap, cp, sc, se = (wb.create_sheet(n) for n in
                              ("Data Quality", "Assumptions", "Account Priorities", "Capacity",
                               "Scenario Comparison"))
    dlog = wb.create_sheet("Decision Log")

    # ---- Assumptions (fixed cell map used by formulas)
    ap["A1"], ap["A1"].font = "Assumptions (blue = editable input)", Font(bold=True, size=13)
    _hdr(ap, 3, ["Coverage hours per account", "customer", "prospect"])
    for r, k, lab in ((4, "high_touch", "High touch"), (5, "light_touch", "Light touch")):
        ap.cell(row=r, column=1, value=lab)
        ap.cell(row=r, column=2, value=a["coverage_hours"][k]["customer"]).font = INPUT
        ap.cell(row=r, column=3, value=a["coverage_hours"][k]["prospect"]).font = INPUT
    _hdr(ap, 7, ["Value capture factor (planning assumption)", "factor"])
    ap["A8"], ap["B8"] = "High touch", a["value_capture"]["high_touch"]
    ap["A9"], ap["B9"] = "Light touch", a["value_capture"]["light_touch"]
    ap["A10"], ap["B10"] = "Not covered", 0
    for c in ("B8", "B9"):
        ap[c].font = INPUT
    rows = [("Capacity buffer (%)", a["capacity_buffer_pct"]), ("High-touch minimum score", a["high_touch_min_score"]),
            ("Light-touch minimum score", a["light_touch_min_score"]),
            ("Max share of hours in high-touch", a["max_high_touch_share"]),
            ("Guardrail: min strategic accounts covered (%)", a["guardrails"]["min_strategic_coverage_pct"]),
            ("Guardrail: max single-region share of value-in-scope (%)", a["guardrails"]["max_single_region_share_pct"]),
            ("Guardrail: min installed-base share (%)", a["guardrails"]["min_customer_share_pct"]),
            ("Guardrail: min new-logo share (%)", a["guardrails"]["min_prospect_share_pct"]),
            ("Capacity shock used in sensitivity (%)", a["capacity_shock_pct"]),
            ("As-of date", a["as_of_date"]), ("Planning period", a["quarter_label"])]
    for i, (lab, val) in enumerate(rows):
        ap.cell(row=12 + i, column=1, value=lab)
        ap.cell(row=12 + i, column=2, value=val).font = INPUT
    _hdr(ap, 25, ["Scenario", "w: value", "w: ICP fit", "w: engagement", "w: strategic",
                  "bonus: customer", "bonus: prospect", "Description"])
    for i, (n, s) in enumerate(a["scenarios"].items()):
        r = 26 + i
        vals = [n, s["weights"]["value"], s["weights"]["icp"], s["weights"]["engagement"], s["weights"]["strategic"],
                s.get("relationship_bonus", {}).get("customer", 0), s.get("relationship_bonus", {}).get("prospect", 0),
                s.get("description", "")]
        for j, v in enumerate(vals, 1):
            c = ap.cell(row=r, column=j, value=v)
            if 2 <= j <= 7:
                c.font = INPUT
    ap["A32"] = ("Note: editing hours or capture factors recalculates Hours and Value-in-scope for the current tier "
                 "assignments. To re-allocate accounts, edit assumptions.json and re-run the planner.")
    ap["A32"].alignment = Alignment(wrap_text=True)
    _widths(ap, [52, 14, 14, 16, 14, 16, 16, 70])

    # ---- Account Priorities
    _hdr(cp, 1, ["Account ID", "Account", "Team", "Region", "Segment", "Relationship", "Strategic", "Value at stake (USD)",
                 "Priority score", "Coverage tier", "Planned hours", "Value-in-scope (USD)", "Ring-fenced", "Reason"])
    ordered = sorted(plan, key=lambda p: (p["team"], p["rank_in_team"]))
    for i, p in enumerate(ordered):
        r = i + 2
        vals = [p["account_id"], p["account_name"], p["team"], p["region"], p["segment"], p["relationship"],
                "Y" if p["strategic"] else "N", round(p["value_at_stake"], 2), p["score"], p["tier"]]
        for j, v in enumerate(vals, 1):
            cp.cell(row=r, column=j, value=v)
        cp.cell(row=r, column=11, value=(f'=IF(J{r}="High touch",IF(F{r}="customer",Assumptions!$B$4,Assumptions!$C$4),'
                                         f'IF(J{r}="Light touch",IF(F{r}="customer",Assumptions!$B$5,Assumptions!$C$5),0))'))
        cp.cell(row=r, column=12, value=(f'=H{r}*IF(J{r}="High touch",Assumptions!$B$8,'
                                         f'IF(J{r}="Light touch",Assumptions!$B$9,0))'))
        cp.cell(row=r, column=13, value="Y" if p["ring"] else "N")
        cp.cell(row=r, column=14, value=p["reason"])
        cp.cell(row=r, column=8).number_format = cp.cell(row=r, column=12).number_format = "#,##0"
        cp.cell(row=r, column=10).fill = GOOD if p["tier"] == "High touch" else WARN if p["tier"] == "Light touch" else PatternFill()
    last = len(ordered) + 1
    cp.freeze_panes = "C2"
    cp.auto_filter.ref = f"A1:N{last}"
    _widths(cp, [11, 30, 18, 15, 12, 13, 10, 16, 10, 13, 10, 17, 10, 70])

    # ---- Capacity (formulas)
    _hdr(sc, 1, ["Team", "Region", "Reps", "Hours per rep", "Buffer (%)", "Available hours", "Planned hours",
                 "Utilization", "High-touch accounts", "Light-touch accounts", "Not covered", "Unconstrained demand (hrs)",
                 "Demand gap (hrs)"])
    tl = sorted(teams)
    AP = "'Account Priorities'"
    for i, t in enumerate(tl):
        r = i + 2
        sc.cell(row=r, column=1, value=t)
        sc.cell(row=r, column=2, value=teams[t]["region"])
        sc.cell(row=r, column=3, value=teams[t]["reps"]).font = INPUT
        sc.cell(row=r, column=4, value=teams[t]["hours_per_rep"]).font = INPUT
        sc.cell(row=r, column=5, value="=Assumptions!$B$12")
        sc.cell(row=r, column=6, value=f"=C{r}*D{r}*(1-E{r}/100)")
        sc.cell(row=r, column=7, value=f"=SUMIFS({AP}!$K$2:$K${last},{AP}!$C$2:$C${last},A{r})")
        sc.cell(row=r, column=8, value=f"=IF(F{r}=0,0,G{r}/F{r})").number_format = "0.0%"
        sc.cell(row=r, column=9, value=f'=COUNTIFS({AP}!$C$2:$C${last},A{r},{AP}!$J$2:$J${last},"High touch")')
        sc.cell(row=r, column=10, value=f'=COUNTIFS({AP}!$C$2:$C${last},A{r},{AP}!$J$2:$J${last},"Light touch")')
        sc.cell(row=r, column=11, value=f'=COUNTIFS({AP}!$C$2:$C${last},A{r},{AP}!$J$2:$J${last},"Not covered")')
        sc.cell(row=r, column=12, value=round(demand.get(t, 0), 1))
        sc.cell(row=r, column=13, value=f"=L{r}-F{r}")
        for col in (6, 7, 12, 13):
            sc.cell(row=r, column=col).number_format = "#,##0"
    cl = len(tl) + 1
    sc.cell(row=cl + 2, column=1, value="Note: capacity shocks passed on the command line are already reflected in Hours per rep.")
    _widths(sc, [20, 16, 8, 14, 10, 16, 15, 12, 18, 18, 13, 24, 16])

    # ---- Scenario Comparison
    names = list(results)
    scn = se
    _hdr(scn, 1, ["Metric"] + names)
    metric_rows = [
        ("Recommended", lambda n, m: "YES" if n == rec else ""),
        ("Value-in-scope index (USD M)", lambda n, m: round(m["value_in_scope"] / 1e6, 3)),
        ("Value at stake covered (USD M)", lambda n, m: round(m["value_at_stake_covered"] / 1e6, 3)),
        ("High-touch accounts", lambda n, m: m["high_touch"]),
        ("Light-touch accounts", lambda n, m: m["light_touch"]),
        ("Eligible accounts not covered / below threshold", lambda n, m: m["not_covered"]),
        ("Hours used", lambda n, m: round(m["hours_used"])),
        ("Capacity utilization (%)", lambda n, m: round(m["utilization_pct"], 1)),
        ("Strategic accounts covered (%)", lambda n, m: round(m["strategic_coverage_pct"], 1)),
        ("In-flight late-stage pipeline protected (%)", lambda n, m: round(m["inflight_protected_pct"], 1)),
        ("Installed-base share of value-in-scope (%)", lambda n, m: round(m["customer_share_pct"], 1)),
        ("New-logo share of value-in-scope (%)", lambda n, m: round(m["prospect_share_pct"], 1)),
        ("Largest region share (%)", lambda n, m: round(m["max_region_share_pct"], 1)),
        ("Largest region", lambda n, m: m["max_region"]),
        ("Guardrails", lambda n, m: "PASS" if m["guardrails_pass"] else "FAIL: " + ", ".join(m["guardrail_failures"])),
    ]
    for i, (lab, fn) in enumerate(metric_rows):
        r = i + 2
        scn.cell(row=r, column=1, value=lab)
        for j, n in enumerate(names):
            c = scn.cell(row=r, column=2 + j, value=fn(n, results[n]["metrics"]))
            if lab == "Guardrails":
                c.fill = GOOD if c.value == "PASS" else BAD
            if n == rec:
                c.font = Font(bold=True)
    scn.cell(row=len(metric_rows) + 3, column=1,
             value="Value-in-scope is a planning index (value at stake x coverage capture factor), not a revenue forecast.")
    ch = BarChart()
    ch.type, ch.title, ch.y_axis.title = "col", "Value-in-scope index by scenario (USD M)", "USD M"
    ch.add_data(Reference(scn, min_col=2, max_col=1 + len(names), min_row=3, max_row=3), from_rows=True)
    ch.set_categories(Reference(scn, min_col=2, max_col=1 + len(names), min_row=1, max_row=1))
    ch.legend, ch.height, ch.width = None, 8, 18
    scn.add_chart(ch, f"A{len(metric_rows) + 6}")
    _widths(scn, [48] + [30] * len(names))

    # ---- Sensitivity
    se2 = wb.create_sheet("Sensitivity")
    se2["A1"] = f"Capacity shock: -{a['capacity_shock_pct']}% for one team at a time ({rec})"
    se2["A1"].font = Font(bold=True, size=12)
    _hdr(se2, 2, ["Team", "Hours lost", "Value-in-scope change (USD)", "Accounts dropped", "Accounts downgraded"])
    for i, r in enumerate(sorted(sens["shocks"], key=lambda r: r["value_in_scope_delta"])):
        for j, v in enumerate([r["team"], round(r["hours_lost"]), round(r["value_in_scope_delta"]),
                               r["accounts_dropped"], r["accounts_downgraded"]], 1):
            se2.cell(row=3 + i, column=j, value=v)
    b = 5 + len(sens["shocks"])
    se2.cell(row=b, column=1, value=(f"Decision stability: {sens['stability_pct']:.0f}% of covered accounts stay covered in all "
                                     f"{sens['perturbations']} weight perturbations (+/-10 pts per weight).")).font = Font(bold=True)
    _hdr(se2, b + 1, ["Marginal account", "Account ID", "Team", "Dropped in N perturbations", "Out of"])
    for i, r in enumerate(sens["marginal"][:25]):
        for j, v in enumerate([r["account_name"], r["account_id"], r["team"], r["dropped_in"], r["of"]], 1):
            se2.cell(row=b + 2 + i, column=j, value=v)
    _widths(se2, [34, 14, 30, 26, 22])

    # ---- Data Quality
    s = clean["stats"]
    dq["A1"], dq["A1"].font = "Data quality and reconciliation", Font(bold=True, size=13)
    for i, (k, v) in enumerate([("Accounts received", s["accounts_raw"]), ("Accounts planned", s["accounts_included"]),
                                ("Accounts excluded", s["accounts_excluded"]), ("Accounts with no issue", s["accounts_clean"]),
                                ("Issues (High / Medium / Low)", f"{s['issues_high']} / {s['issues_medium']} / {s['issues_low']}"),
                                ("Data confidence", f"{s['confidence_label']} ({s['confidence_pct']}% of received accounts clean)")]):
        dq.cell(row=3 + i, column=1, value=k)
        dq.cell(row=3 + i, column=2, value=v)
    _hdr(dq, 11, ["Severity", "Table", "Record", "Rule", "Detail", "Action taken"])
    sev = {"High": BAD, "Medium": WARN, "Low": PatternFill()}
    for i, it in enumerate(sorted(clean["issues"], key=lambda x: ({"High": 0, "Medium": 1, "Low": 2}[x["severity"]], x["table"], x["record"]))):
        for j, k in enumerate(["severity", "table", "record", "rule", "detail", "action"], 1):
            c = dq.cell(row=12 + i, column=j, value=it[k])
            if j == 1:
                c.fill = sev[it["severity"]]
    _widths(dq, [30, 16, 14, 30, 60, 55])

    # ---- Decision Log
    _hdr(dlog, 1, ["ID", "Date", "Decision requested", "Recommendation", "Rationale", "Key assumptions to validate",
                   "Main risks", "Approver", "Status", "Review date"])
    big = max(sens["shocks"], key=lambda r: -r["value_in_scope_delta"])
    dlog.append(["GTM-001", date.today().isoformat(),
                 f"Approve {a['quarter_label']} coverage plan and capacity allocation",
                 f"Adopt scenario '{rec}': {metrics['high_touch']} high-touch + {metrics['light_touch']} light-touch accounts",
                 ("Highest value-in-scope among scenarios that satisfy all guardrails." if metrics["guardrails_pass"]
                  else "No scenario satisfies every guardrail; chosen for best strategic coverage. Needs explicit executive trade-off."),
                 "Hours per tier, value capture factors, ICP/engagement inputs, rep capacity after buffer",
                 f"Largest exposure: -{a['capacity_shock_pct']}% capacity in {big['team']} ({big['accounts_dropped']} accounts dropped)",
                 "", "Proposed", ""])
    _widths(dlog, [10, 12, 40, 50, 55, 50, 55, 14, 12, 12])

    # ---- Executive Summary with reconciliation
    ws["A1"], ws["A1"].font = f"Quarterly GTM Investment Plan: {a['quarter_label']}", Font(bold=True, size=15)
    ws["A2"] = "Demo data is fictional. Planning model, not a revenue forecast."
    ws["A3"], ws["B3"] = "Recommended scenario", rec
    ws["A4"], ws["B4"] = "Guardrails", "PASS" if metrics["guardrails_pass"] else "FAIL: " + ", ".join(metrics["guardrail_failures"])
    ws["A5"], ws["B5"] = "Data confidence", f"{clean['stats']['confidence_label']} ({clean['stats']['confidence_pct']}%)"
    _hdr(ws, 7, ["Reconciliation: Excel formulas vs. Python engine", "Excel (formula)", "Engine", "Check"])
    items = [("Planned hours", f"=SUM(Capacity!G2:G{cl})", round(metrics["hours_used"], 2)),
             ("Available hours", f"=SUM(Capacity!F2:F{cl})", round(metrics["hours_available"], 2)),
             ("High-touch accounts", f'=COUNTIF({AP}!J2:J{last},"High touch")', metrics["high_touch"]),
             ("Light-touch accounts", f'=COUNTIF({AP}!J2:J{last},"Light touch")', metrics["light_touch"]),
             ("Value-in-scope (USD)", f"=SUM({AP}!L2:L{last})", round(metrics["value_in_scope"], 2))]
    for i, (lab, f, v) in enumerate(items):
        r = 8 + i
        ws.cell(row=r, column=1, value=lab)
        ws.cell(row=r, column=2, value=f).number_format = "#,##0"
        ws.cell(row=r, column=3, value=v).number_format = "#,##0"
        ws.cell(row=r, column=4, value=f'=IF(ABS(B{r}-C{r})<1,"OK","CHECK")')
    img = XLImage(str(ctx["charts"]["scenarios"]))
    img.width, img.height = 720, 336
    ws.add_image(img, "A15")
    _widths(ws, [52, 26, 22, 12])
    wb.save(path)
